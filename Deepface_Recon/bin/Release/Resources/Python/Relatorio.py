import os

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches
from docx.shared import RGBColor
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pandas as pd
from datetime import datetime

from numpy.polynomial.polyutils import trimseq


def DocxCreation(LogoPath: str, CsvTempsPath: str, TimelogPath: str, PathToSave: str):
    Data_Hora = datetime.strftime(datetime.now(), '%Y-%m-%d %H.%M.%S')

    Relatório = Document()#Criação inicial do documento

    img = Relatório.add_picture(LogoPath, height=Inches(3.334646))#Adicionar logo no começo
    ultimo = Relatório.paragraphs[-1]#Selecionar paragrafo mais recente
    ultimo.alignment = WD_ALIGN_PARAGRAPH.CENTER#Centralizar tudo nele para centralizar a imagem

    Titulo = Relatório.add_paragraph()#Criação do texto do título
    TituloFormat = Titulo.paragraph_format  # Fazendo texto ser formatavel
    TituloFormat.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER  # Centralizando o texto
    TituloRun = Titulo.add_run("Relatório de Gravação\n (" + Data_Hora + ")")#Texto do Titulo
    TituloRun.font.size = Pt(28)#Tamanho da fonte

    Descrição = Relatório.add_paragraph()
    DescriçãoFormat = Descrição.paragraph_format
    DescriçãoFormat.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    DescriçãoRun = Descrição.add_run("Documento representa a documentação de todas as pessoas registradas que foram capturadas pelo “Águia Eye”")
    DescriçãoRun.font.size = Pt(11)

    skeleton = {"Confidence": [], "Candidate": [], "Database": [], "TimeFrame": []} #Esqueleto sobre como os dados a serem importados pelos arquivos csv vão ser organizados
    DataCompleta = pd.DataFrame(skeleton)#Criando um DataFrame com o dito esqueleto

    log = pd.read_csv(TimelogPath)#Lendo um arquvivo csv que eventualmente vai ter seus dados em "DataCompleta"

    for cvs in os.listdir(CsvTempsPath):#Procurando outros arquivos csv que entraram em "DataCompleta"
        if cvs != None:
            FullCvs = os.path.join(CsvTempsPath, cvs)
            Dataframe = pd.read_csv(FullCvs)#Lendo arquvivos csv que eventualmente vão ter seus dados em "DataCompleta"
            for index, row in log.iterrows():#Juntando todas as informações em "DataCompleta"
                if index != None:
                    if row["FileName"] == Dataframe.loc[0,'Candidate']:
                        temp = pd.DataFrame(skeleton)
                        temp.loc[0,'Confidence'] = Dataframe.loc[0, "Confidence"]
                        temp.loc[0, 'Candidate'] = Dataframe.loc[0, "Candidate"]
                        temp.loc[0, 'Database'] = Dataframe.loc[0, "Database"]
                        temp.loc[0, 'TimeFrame'] = row['TakenTime']
                        DataCompleta =  pd.concat([DataCompleta, temp], ignore_index=True) #Todas as informações necessarias finalmente em "DataCompleta"

    PreviousEntry =  None
    for index, row in DataCompleta.iterrows():
        if(row["Database"] != PreviousEntry):#Verificando se é a mesma pessoa que apareceu na última imagem
            Relatório.add_page_break()
            PreviousEntry = row["Database"]
            EntryList = DataCompleta[DataCompleta['Database'] == row['Database']]
            DbString = str(row['Database'])#Usamos a database para conseguir o nome da imagem
            Dblist = os.path.split(DbString)
            DbName = Dblist[len(Dblist) - 1]
            DbName.strip(")")

            pic = Relatório.add_picture(row["Candidate"].replace("PicturesTS", "PicturesTA"), height=Inches(3.93701)) #Adiciona a primeira foto qual o individuo foi reconhecido
            ultimo = Relatório.paragraphs[-1]
            ultimo.alignment = WD_ALIGN_PARAGRAPH.CENTER#Centraliza a imagem

            Documentado = Relatório.add_paragraph()
            DocumentadoFormat = Documentado.paragraph_format
            DocumentadoFormat.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            DocumentadoRun = Documentado.add_run("DOCUMENTADO PELA CÂMERA AS:")
            DocumentadoRun.font.size = Pt(12)

            Identificação = Relatório.add_paragraph()
            IdentificaçãoFormat = Identificação.paragraph_format
            IdentificaçãoFormat.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            IdentificaçãoRun = Identificação.add_run("("+ str(row['TimeFrame']) + ")")#Usando os dados da coluna TimeFrame para saber o tempo qual a imagem foi tirada
            IdentificaçãoRun.font.size = Pt(12)
            IdentificaçãoRun.font.color.rgb = RGBColor(255, 0, 0)

            Individuo = Relatório.add_paragraph()
            IndividuoFormat = Individuo.paragraph_format
            IndividuoFormat.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            IndividuoRun = Individuo.add_run("INDIVÍDUO:")
            IndividuoRun.font.size = Pt(12)

            Pessoa = Relatório.add_paragraph()
            PessoaFormat = Pessoa.paragraph_format
            PessoaFormat.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

            PessoaRun = Pessoa.add_run(DbName)#Usando o nome que adquirimos a partir da database
            PessoaRun.font.size = Pt(12)
            PessoaRun.font.color.rgb = RGBColor(255, 0, 0)

            Quantidade = Relatório.add_paragraph()
            QuantidadeFormat = Quantidade.paragraph_format
            QuantidadeFormat.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            QuantidadeRun = Quantidade.add_run("QUANTIDADES DE VEZES QUE APARECE:")
            QuantidadeRun.font.size = Pt(12)

            Repetição = Relatório.add_paragraph()
            RepetiçãoFormat = Repetição.paragraph_format
            RepetiçãoFormat.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            RepetiçãoRun = Repetição.add_run("(" + str(len(EntryList)) + ")")
            RepetiçãoRun.font.size = Pt(12)
            RepetiçãoRun.font.color.rgb = RGBColor(255, 0, 0)
    Data_Hora = datetime.strftime(datetime.now(), '%Y-%m-%d %H.%M.%S')#Hora e dia atual
    Relatório.save(PathToSave + r"\Relatorio " + str(Data_Hora) + ".docx")#Salvar o relatório
    #os.remove(r"C:\Users\Luciano Fonte\Desktop\a.docx")

#all of these will need to be handled by Pre2Água-Eye.py, so now remember to get the path of the timelog folder there, also the logo too.
#DocxCreation(r"C:\Users\Luciano Fonte\Desktop\Deepface Prj\Builds\FullApps\Build2\Resources\Relatórios\re\Logo.png",
             #r"C:\Users\Luciano Fonte\Desktop\Deepface Prj\Builds\FullApps\Build2\Resources\Sessions\4195\CSVtemps",
             #r"C:\Users\Luciano Fonte\Desktop\Deepface Prj\Builds\FullApps\Build2\Resources\Sessions\4195\TimeLog\TimeDecoder.csv")


